"""Word documents (.docx)."""

from __future__ import annotations

from pathlib import Path

from ..core import ReadOptions, ReadResult
from ..registry import reader
from ._doctools import chunk_paragraphs, doc_metadata


@reader("docx", category="Documents", extensions=(".docx",), requires=("docx",), extra="docs")
def read_docx(path: Path, opts: ReadOptions) -> ReadResult:
    import docx

    d = docx.Document(str(path))
    paras, headings, title = [], [], ""
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style == "title" and not title:
            title = text
        elif style.startswith("heading"):
            headings.append(text)
        paras.append(text)
    tables = [
        {"page": None, "rows": [[c.text for c in r.cells] for r in t.rows[:30]]}
        for t in d.tables[:15]
    ]
    cp = d.core_properties
    authors = [a.strip() for a in (cp.author or "").split(",") if a.strip()]
    pages = chunk_paragraphs(paras)
    meta = doc_metadata(
        pages, title=title or (cp.title or ""), authors=authors, sections=headings, tables=tables
    )
    meta["page_note"] = "DOCX has no fixed pages: 'page' = ~500-word part"
    return ReadResult("document", None, meta, pages)
