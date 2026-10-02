"""Generate tiny synthetic sample files for every supported format.

Each entry: name -> (path, expected reader, expected kind). Formats whose writer library
is missing are skipped (and so are their tests).
"""

from __future__ import annotations

import contextlib
import gzip
import json
import shutil
import tarfile
import zipfile
from collections.abc import Callable
from pathlib import Path

import numpy as np
import pandas as pd
import xarray as xr

from nasa_explorer._native import preload_before_eccodes

preload_before_eccodes()  # see nasa_explorer/_native.py

RNG = np.random.default_rng(42)
TIMES = pd.date_range("2001-01-01", periods=48, freq="MS")
LAT = np.arange(20.0, 41.0, 2.5)
LON = np.arange(30.0, 51.0, 2.5)

Sample = tuple[Path, str, str]


def grid_ds() -> xr.Dataset:
    trend = 0.05 * np.arange(TIMES.size)[:, None, None]
    data = 290 + trend + RNG.normal(0, 0.3, (TIMES.size, LAT.size, LON.size))
    return xr.Dataset(
        {
            "t2m": (
                ("time", "lat", "lon"),
                data,
                {"units": "K", "long_name": "2 m air temperature"},
            ),
            "precip": (("time", "lat", "lon"), RNG.gamma(2, 1, data.shape), {"units": "mm/day"}),
        },
        coords={
            "time": TIMES,
            "lat": ("lat", LAT, {"units": "degrees_north"}),
            "lon": ("lon", LON, {"units": "degrees_east"}),
        },
        attrs={"title": "synthetic test grid", "source": "nasa-data-explorer tests"},
    )


def table_df() -> pd.DataFrame:
    n = 60
    return pd.DataFrame(
        {
            "date": pd.date_range("2020-01-01", periods=n, freq="D"),
            "latitude": RNG.uniform(29, 33, n),
            "longitude": RNG.uniform(35, 39, n),
            "brightness": 300 + np.arange(n) * 0.5 + RNG.normal(0, 1, n),
            "frp": RNG.gamma(2, 5, n),
            "confidence": RNG.integers(0, 100, n),
        }
    )


def _tabular(d: Path) -> dict[str, Sample]:
    df = table_df()
    out = {}
    df.to_csv(d / "fires.csv", index=False)
    out["csv"] = (d / "fires.csv", "delimited-text", "table")
    df.to_csv(d / "fires.tsv", sep="\t", index=False)
    out["tsv"] = (d / "fires.tsv", "delimited-text", "table")
    df.to_csv(d / "semicolon.txt", sep=";", index=False)
    out["txt_delimited"] = (d / "semicolon.txt", "delimited-text", "table")
    rows = "\n".join(
        f"{t.year},{t.month},{t.day},{v:.2f}"
        for t, v in zip(
            pd.date_range("2010-01-01", periods=40, freq="D"), RNG.normal(20, 2, 40), strict=True
        )
    )
    (d / "power_like.csv").write_text(
        "-BEGIN HEADER-\nNASA/POWER Source Native Resolution Daily Data\nT2M  MERRA-2 Temperature at 2 Meters (C)\n"
        "-END HEADER-\nYEAR,MO,DY,T2M\n" + rows + "\n2010,2,10,-999\n"
    )
    out["power_header_csv"] = (d / "power_like.csv", "delimited-text", "table")
    (d / "whitespace.dat").write_text(
        "x y z\n" + "\n".join(f"{i} {i * 2.5:.1f} {RNG.normal():.3f}" for i in range(30))
    )
    out["whitespace_dat"] = (d / "whitespace.dat", "delimited-text", "table")
    df.to_excel(d / "book.xlsx", index=False)
    out["xlsx"] = (d / "book.xlsx", "excel", "table")
    (d / "records.json").write_text(df.head(20).to_json(orient="records", date_format="iso"))
    out["json_records"] = (d / "records.json", "json", "table")
    days = pd.date_range("2015-01-01", periods=30, freq="D").strftime("%Y%m%d")
    power = {
        "type": "Feature",
        "geometry": {"type": "Point", "coordinates": [35.9, 31.9]},
        "properties": {
            "parameter": {
                "T2M": {k: float(v) for k, v in zip(days, RNG.normal(10, 3, 30), strict=True)},
                "RH2M": {k: float(v) for k, v in zip(days, RNG.uniform(20, 80, 30), strict=True)},
            }
        },
    }
    (d / "power.json").write_text(json.dumps(power))
    out["json_nested"] = (d / "power.json", "json", "table")
    df.to_parquet(d / "fires.parquet")
    out["parquet"] = (d / "fires.parquet", "parquet", "table")
    return out


def _gridded(d: Path) -> dict[str, Sample]:
    ds = grid_ds()
    out = {}
    ds.to_netcdf(d / "grid.nc", engine="netcdf4")
    out["netcdf4"] = (d / "grid.nc", "netcdf", "grid")
    ds.to_netcdf(d / "classic.nc", engine="scipy", format="NETCDF3_64BIT")
    out["netcdf3"] = (d / "classic.nc", "netcdf", "grid")
    packed = ds[["t2m"]].copy()
    packed["t2m"].values[0, 0, 0] = np.nan
    packed.to_netcdf(
        d / "packed.nc",
        encoding={
            "t2m": {
                "dtype": "int16",
                "scale_factor": 0.01,
                "add_offset": 290.0,
                "_FillValue": -32767,
            }
        },
    )
    out["netcdf_packed"] = (d / "packed.nc", "netcdf", "grid")
    ds.to_netcdf(d / "grid.h5", engine="h5netcdf")
    out["hdf5_xarray"] = (d / "grid.h5", "hdf5", "grid")
    shutil.copy(d / "grid.nc", d / "dataset_without_extension")
    out["no_extension"] = (d / "dataset_without_extension", "hdf5", "grid")
    shutil.copy(d / "classic.nc", d / "mislabelled.csv")
    out["wrong_extension"] = (d / "mislabelled.csv", "netcdf", "grid")

    import h5py

    with h5py.File(d / "nisar_like.h5", "w") as f:
        f.attrs["mission_name"] = "NISAR-like synthetic"
        g = f.create_group("science/LSAR/GCOV/grids/frequencyA")
        raw = RNG.integers(0, 1000, (64, 64)).astype("int16")
        raw[:4, :4] = -9999
        dset = g.create_dataset("HHHH", data=raw)
        dset.attrs.update(
            {
                "_FillValue": np.int16(-9999),
                "scale_factor": 0.001,
                "units": "1",
                "long_name": "HH backscatter",
            }
        )
        g.create_dataset("xCoordinates", data=np.linspace(500000, 510000, 64))
        f.create_group("metadata/orbit").create_dataset("time", data=np.arange(10.0))
    out["hdf5_tree"] = (d / "nisar_like.h5", "hdf5", "tree")

    with contextlib.suppress(ImportError):
        from pyhdf.SD import SD, SDC

        sd = SD(str(d / "modis_like.hdf"), SDC.WRITE | SDC.CREATE)
        raw = RNG.integers(0, 10000, (20, 30)).astype("int16")
        raw[0, :5] = -3000
        sds = sd.create("NDVI", SDC.INT16, raw.shape)
        sds.dim(0).setname("YDim")
        sds.dim(1).setname("XDim")
        sds[:] = raw
        sds.setfillvalue(-3000)
        sds.scale_factor, sds.add_offset = 0.0001, 0.0
        sds.units = "NDVI"
        sds.endaccess()
        sd.end()
        out["hdf4"] = (d / "modis_like.hdf", "hdf4", "grid")

    with contextlib.suppress(ImportError):
        import eccodes

        with open(d / "sample.grib2", "wb") as fh:
            for step in range(3):
                gid = eccodes.codes_grib_new_from_samples("regular_ll_sfc_grib2")
                eccodes.codes_set(gid, "dataDate", 20240101 + step)
                eccodes.codes_set_values(gid, 280 + RNG.normal(0, 2, 16 * 31))
                eccodes.codes_write(gid, fh)
                eccodes.codes_release(gid)
        out["grib"] = (d / "sample.grib2", "grib", "grid")

    with contextlib.suppress(ImportError):
        import zarr  # noqa: F401

        ds.to_zarr(d / "store.zarr", mode="w")
        out["zarr"] = (d / "store.zarr", "zarr", "grid")
    return out


def _geo(d: Path) -> dict[str, Sample]:
    out = {}
    try:
        import geopandas as gpd
        import rasterio
        from rasterio.transform import from_origin
        from shapely.geometry import Point, box
    except ImportError:
        return out
    dem = (RNG.random((50, 60)) * 1000).astype("float32")
    with rasterio.open(
        d / "dem.tif",
        "w",
        driver="GTiff",
        height=50,
        width=60,
        count=1,
        dtype="float32",
        crs="EPSG:4326",
        transform=from_origin(35.0, 33.0, 0.05, 0.05),
        nodata=-9999,
    ) as dst:
        dst.write(dem, 1)
    out["geotiff"] = (d / "dem.tif", "geotiff", "grid")
    pts = gpd.GeoDataFrame(
        {"name": [f"site{i}" for i in range(10)], "value": RNG.normal(size=10)},
        geometry=[Point(35 + i * 0.1, 31 + i * 0.05) for i in range(10)],
        crs="EPSG:4326",
    )
    pts.to_file(d / "points.shp")
    out["shapefile"] = (d / "points.shp", "vector", "table")
    polys = gpd.GeoDataFrame(
        {"zone": ["a", "b"], "area_km2": [1.0, 2.0]},
        geometry=[box(35, 31, 36, 32), box(36, 31, 37, 32)],
        crs="EPSG:4326",
    )
    polys.to_file(d / "zones.geojson", driver="GeoJSON")
    out["geojson"] = (d / "zones.geojson", "vector", "table")
    pts.to_file(d / "sites.gpkg", layer="sites", driver="GPKG")
    out["gpkg"] = (d / "sites.gpkg", "vector", "table")
    pts[["name", "geometry"]].to_file(d / "sites.kml", driver="KML")
    out["kml"] = (d / "sites.kml", "vector", "table")
    with zipfile.ZipFile(d / "sites.kmz", "w") as zf:
        zf.write(d / "sites.kml", "doc.kml")
    out["kmz"] = (d / "sites.kmz", "vector", "table")
    return out


def _astro(d: Path) -> dict[str, Sample]:
    out = {}
    try:
        from astropy.io import fits
        from astropy.table import Table
    except ImportError:
        return out
    img = RNG.normal(100, 10, (40, 50)).astype("float32")
    hdr = fits.Header(
        {
            "OBJECT": "M31",
            "TELESCOP": "HST",
            "BUNIT": "electrons/s",
            "CTYPE1": "RA---TAN",
            "CTYPE2": "DEC--TAN",
            "CRVAL1": 10.68,
            "CRVAL2": 41.27,
            "CRPIX1": 25,
            "CRPIX2": 20,
            "CDELT1": -0.0001,
            "CDELT2": 0.0001,
        }
    )
    cat = Table(
        {
            "ra": RNG.uniform(10, 11, 30),
            "dec": RNG.uniform(41, 42, 30),
            "mag": RNG.normal(20, 1, 30),
        }
    )
    fits.HDUList(
        [
            fits.PrimaryHDU(img, header=hdr),
            fits.ImageHDU(img * 2, name="SCI2"),
            fits.BinTableHDU(cat, name="CATALOG"),
        ]
    ).writeto(d / "multi.fits", overwrite=True)
    out["fits_image"] = (d / "multi.fits", "fits", "grid")
    fits.HDUList([fits.PrimaryHDU(), fits.BinTableHDU(cat, name="CATALOG")]).writeto(
        d / "catalog.fits", overwrite=True
    )
    out["fits_table"] = (d / "catalog.fits", "fits", "table")
    return out


def _images(d: Path) -> dict[str, Sample]:
    from PIL import Image

    arr = (RNG.random((40, 60, 3)) * 255).astype("uint8")
    Image.fromarray(arr).save(d / "photo.png")
    Image.fromarray(arr).save(d / "photo.jpg")
    Image.fromarray(arr).save(d / "plain.tif")
    return {
        "png": (d / "photo.png", "image", "image"),
        "jpg": (d / "photo.jpg", "image", "image"),
        "tiff_plain": (d / "plain.tif", "image", "image"),
    }


PAPER = [
    "Warming Trends over the Levant from MODIS Land Surface Temperature\n"
    "Lina Haddad, Omar Saleh and Maya Khoury\n\nAbstract\nWe analyse MODIS Terra LST from 2001 to 2020 and find "
    "a warming of 0.45 K per decade (p < 0.01).\n\n1 Introduction\nHeat extremes are increasing in the region.",
    "2 Data and Methods\nWe used MODIS MOD11A2 and GPM IMERG precipitation, validated against GLDAS.\n"
    "Figure 1: Mean land surface temperature anomaly for 2001-2020.\n\n3 Results\nThe trend is 0.45 K per decade.",
    "4 Conclusion\nUrban areas warm faster. Limitations include cloud gaps.\nTable 1: Station list.\n\nReferences\n[1] Wan, Z. (2014).",
]


def _documents(d: Path) -> dict[str, Sample]:
    out = {}
    with contextlib.suppress(ImportError):
        import pymupdf

        doc = pymupdf.open()
        for text in PAPER:
            page = doc.new_page()
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
        doc.set_metadata(
            {
                "title": "Warming Trends over the Levant from MODIS Land Surface Temperature",
                "author": "Lina Haddad, Omar Saleh, Maya Khoury",
            }
        )
        doc.save(d / "paper.pdf")
        out["pdf"] = (d / "paper.pdf", "pdf", "document")

        from PIL import Image, ImageDraw, ImageFont

        img = Image.new("L", (1200, 400), 255)
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("DejaVuSans.ttf", 40)
        except OSError:
            font = ImageFont.load_default(size=40)
        draw.text((40, 60), "SCANNED REPORT", fill=0, font=font)
        draw.text((40, 160), "Landsat surface reflectance", fill=0, font=font)
        img.save(d / "scan.png")
        scan = pymupdf.open()
        page = scan.new_page(width=600, height=200)
        page.insert_image(page.rect, filename=str(d / "scan.png"))
        scan.save(d / "scanned.pdf")
        (d / "scan.png").unlink()
        out["pdf_scanned"] = (d / "scanned.pdf", "pdf", "document")
    with contextlib.suppress(ImportError):
        import docx

        doc = docx.Document()
        doc.add_heading("SMAP Soil Moisture Notes", 0)
        doc.add_heading("Introduction", 1)
        doc.add_paragraph("SMAP and GRACE observations were compared across Jordan.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = (
            "site",
            "value",
            "Amman",
            "0.21",
        )
        doc.core_properties.author = "Test Author"
        doc.save(d / "notes.docx")
        out["docx"] = (d / "notes.docx", "docx", "document")
    (d / "page.html").write_text(
        "<!doctype html><html><head><title>VIIRS fire page</title><meta name='author' content='Web Team'></head><body>"
        "<h1>VIIRS Active Fires</h1><p>FIRMS distributes VIIRS detections.</p>"
        "<figure><img src='x.png'><figcaption>Figure 2: Fire counts</figcaption></figure>"
        "<table><tr><th>day</th><th>fires</th></tr><tr><td>1</td><td>12</td></tr></table></body></html>"
    )
    out["html"] = (d / "page.html", "html", "document")
    (d / "readme.md").write_text(
        "# GEDI Canopy Height\n\n## Data\n\nGEDI L2A footprints.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    )
    out["markdown"] = (d / "readme.md", "markdown", "document")
    (d / "notes.txt").write_text(
        "These are free-form field notes about ICESat-2 tracks.\nNothing tabular here at all, "
        "just sentences that describe the campaign.\nAnother line of prose follows.\n"
    )
    out["txt_prose"] = (d / "notes.txt", "text", "document")
    return out


def _misc(d: Path, made: dict[str, Sample]) -> dict[str, Sample]:
    out = {}
    (d / "mystery.bin").write_bytes(
        b"\x00\x01\x02NOTAFORMAT" + RNG.integers(0, 255, 200, dtype=np.uint8).tobytes()
    )
    out["unknown"] = (d / "mystery.bin", "unknown", "binary")
    with zipfile.ZipFile(d / "bundle.zip", "w") as zf:
        zf.write(made["csv"][0], "inner/fires.csv")
        zf.write(made["netcdf4"][0], "inner/grid.nc")
    out["zip"] = (d / "bundle.zip", "archive", "archive")
    with open(made["csv"][0], "rb") as src, gzip.open(d / "fires.csv.gz", "wb") as dst:
        dst.write(src.read())
    out["gz"] = (d / "fires.csv.gz", "archive", "archive")
    with tarfile.open(d / "bundle.tar.gz", "w:gz") as tf:
        tf.add(made["csv"][0], "fires.csv")
        tf.add(made["netcdf3"][0], "classic.nc")
    out["targz"] = (d / "bundle.tar.gz", "archive", "archive")
    return out


def make_samples(d: Path) -> dict[str, Sample]:
    d.mkdir(parents=True, exist_ok=True)
    made: dict[str, Sample] = {}
    builders: list[Callable[[Path], dict[str, Sample]]] = [
        _tabular,
        _gridded,
        _geo,
        _astro,
        _images,
        _documents,
    ]
    for build in builders:
        made.update(build(d))
    made.update(_misc(d, made))
    return made
